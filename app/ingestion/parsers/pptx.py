from io import BytesIO
from pptx import Presentation
from .base import DocumentParser

class PptxParser(DocumentParser):
    def supports(self, filename, content_type=None):
        return filename.lower().endswith('.pptx')

    def parse(self, data, filename):
        prs = Presentation(BytesIO(data))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            text = "\n".join(
                shape.text for shape in slide.shapes
                if hasattr(shape, 'text') and shape.text.strip()
            )
            slides.append(f"[Slide {i}]\n{text}")
        return "\n\n".join(slides), {'parser': 'pptx', 'slides': len(prs.slides)}
