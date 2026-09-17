from io import BytesIO
from docx import Document
from .base import DocumentParser

class DocxParser(DocumentParser):
    def supports(self, filename, content_type=None):
        return filename.lower().endswith('.docx')

    def parse(self, data, filename):
        doc = Document(BytesIO(data))
        blocks = []
        for p in doc.paragraphs:
            if p.text.strip():
                blocks.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                blocks.append(' | '.join(cell.text for cell in row.cells))
        return "\n".join(blocks), {'parser': 'docx'}
