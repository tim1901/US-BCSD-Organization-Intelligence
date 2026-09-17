from io import BytesIO
from openpyxl import load_workbook
from .base import DocumentParser

class XlsxParser(DocumentParser):
    def supports(self, filename, content_type=None):
        return filename.lower().endswith('.xlsx')

    def parse(self, data, filename):
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[Sheet {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) if v is not None else '' for v in row]
                parts.append(' | '.join(vals))
        return "\n".join(parts), {'parser': 'xlsx', 'sheets': len(wb.worksheets)}
